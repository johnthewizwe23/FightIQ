"""Build the FightIQ Excel workbook.

Sheets:
  Guide          how to use it, how the model works, key results
  Predictor      enter a card + odds; Excel formulas run the model (no Python needed)
  Tracking       log picks and bets; record, units and ROI calculate themselves
  Fighters       every UFC fighter's current stats (as of the latest event in the data)
  Model          coefficients, blend weights, bet rules and defaults used by the Predictor
  Backtest       walk-forward results, 2015 onwards
  Fight Database every bout with both fighters' stats as they stood before it
"""
from __future__ import annotations

import shutil
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config, data, features
from . import evaluate as ev
from . import odds as odds_mod
from . import tracking as tracking_mod
from .model import FightModel

FONT = "Arial"
BLUE = "0000FF"
GREEN = "008000"
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
HELPER_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
PREDICTOR_ROWS = 20

FEATURE_TEXT = {
    "slpm": "Sig. strikes landed per minute (A − B)",
    "exp_str_acc": "Expected strike accuracy: √(acc × (1 − opp. defence)), A − B",
    "sapm": "Sig. strikes absorbed per minute (A − B)",
    "td_avg": "Takedowns landed per 15 min (A − B)",
    "reach": "Reach in inches (A − B)",
    "age": "Age in years on fight night (A − B)",
    "elo": "Elo rating ÷ 100 (A − B)",
    "ctrl_avg": "Control time, minutes per 15 (A − B)",
    "experience": "ln(1 + UFC fights), A − B",
    "layoff": "ln(1 + days since last fight), A − B",
    "age_old": "Years past age 33 (A − B)",
    "debut": "UFC debut flag (A − B)",
}


def _font(**kw) -> Font:
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def _header(ws, row: int, labels: list[str], start_col: int = 1, fill=HEADER_FILL, color="FFFFFF") -> None:
    for j, label in enumerate(labels, start=start_col):
        c = ws.cell(row, j, label)
        c.font = _font(bold=True, color=color)
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)


def _widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def build_workbook(path=None, card: pd.DataFrame | None = None, event_name: str = "", event_date=None,
                   live_odds: bool = False) -> str:
    """Write the workbook.

    With `card` (fighter_a, fighter_b, optional odds_a/odds_b) the Predictor gets that card; without one
    it keeps the card already in the existing workbook. `live_odds` fills the odds from The Odds API.
    Tracking: rows already there are kept, every bout with odds (the previous card and this one) is
    logged, and results are filled in from the data. The old file is copied to output/backup/ first.
    """
    path = Path(path or config.OUTPUT_DIR / "FightIQ_Model.xlsx")
    tracking_rows = _read_tracking(path)
    previous = _read_predictor(path)
    if card is None and previous is not None:
        card, event_name, event_date = previous
    if live_odds and card is not None:
        card = odds_mod.add_live_odds(card)
    model = FightModel.load(config.MODEL_PATH)
    priors = model.meta["priors"]

    fights, rounds = data.load_fights()
    fighters = data.load_fighters()
    last_event = fights["date"].max()
    as_of = last_event + pd.Timedelta(days=1)
    X = features.build(fights, rounds, fighters, priors=priors)
    X = ev.attach_odds(X, data.load_odds())
    print("running walk-forward backtest for the Backtest sheet...")
    P = ev.walk_forward_blend(ev.walk_forward(X, model.features, model.C))
    profiles = _fighter_table(fights, rounds, fighters, priors, as_of)
    if card is not None:
        card = _card_display_names(card, fights, fighters, profiles)
    tracking_rows = _update_tracking(tracking_rows, previous, (card, event_name, event_date), fights)

    wb = Workbook()
    wb._named_styles["Normal"].font = _font()
    ws_guide = wb.active
    ws_guide.title = "Guide"
    ws_pred = wb.create_sheet("Predictor")
    ws_track = wb.create_sheet("Tracking")
    ws_f = wb.create_sheet("Fighters")
    ws_m = wb.create_sheet("Model")
    ws_bt = wb.create_sheet("Backtest")
    ws_db = wb.create_sheet("Fight Database")

    cells = _model_sheet(ws_m, model, priors)
    fcols = _fighters_sheet(ws_f, profiles, last_event)
    _predictor_sheet(ws_pred, model, cells, fcols, len(profiles), profiles, card, event_name, event_date)
    _tracking_sheet(ws_track, tracking_rows)
    summary = _backtest_sheet(ws_bt, P)
    _database_sheet(ws_db, X, P, model)
    _guide_sheet(ws_guide, model, summary, last_event, len(fights))

    wb.calculation.fullCalcOnLoad = True  # openpyxl stores no results; Excel computes them on open
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        backup = path.parent / "backup" / f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}"
        backup.parent.mkdir(exist_ok=True)
        shutil.copy2(path, backup)
    try:
        wb.save(path)
    except PermissionError as exc:
        raise SystemExit(f"can't write {path}: close it in Excel and run the command again") from exc
    return str(path)


TRACKING_TOP = 11  # first data row of the Tracking sheet
TRACKING_INPUT_COLS = 10  # columns A-J are typed in; K-L are formulas


def _read_tracking(path: Path) -> list[list]:
    """Rows the user has typed into an existing workbook's Tracking sheet."""
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, read_only=True)
    except Exception:
        return []
    if "Tracking" not in wb.sheetnames:
        return []
    rows = []
    for row in wb["Tracking"].iter_rows(min_row=TRACKING_TOP, max_col=TRACKING_INPUT_COLS, values_only=True):
        if any(v not in (None, "") for v in row):
            rows.append(list(row))
    wb.close()
    return rows


def _read_predictor(path: Path):
    """(card, event name, event date) from an existing workbook's Predictor, or None."""
    if not path.exists():
        return None
    try:
        wb = load_workbook(path)
        cached = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return None
    if "Predictor" not in wb.sheetnames:
        return None
    ws = wb["Predictor"]
    title = str(ws["A1"].value or "")
    event = title.split(" - ", 1)[1] if " - " in title else ""
    when = ws["B3"].value
    if not isinstance(when, datetime):
        when = cached["Predictor"]["B3"].value
    when = when if isinstance(when, datetime) else datetime.combine(date.today(), datetime.min.time())
    rows = []
    for r in range(6, 6 + PREDICTOR_ROWS):
        a, b, oa, ob = (ws.cell(r, c).value for c in range(1, 5))
        if a and b and str(a).strip() and str(b).strip():
            num = _odds_value
            rows.append({"fighter_a": str(a).strip(), "fighter_b": str(b).strip(), "odds_a": num(oa), "odds_b": num(ob)})
    cached.close()
    if not rows:
        return None
    return pd.DataFrame(rows), event, when


def _odds_value(v) -> float:
    """Odds typed as a number or as text ('1.85', '1,85', '+150', '-200')."""
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and v.strip():
        from .card import parse_odds
        try:
            return parse_odds(v.strip().replace(",", "."))
        except ValueError:
            return np.nan
    return np.nan


def _update_tracking(rows: list[list], previous, current, fights) -> list[list]:
    """Log the previous and current Predictor cards (bouts with odds) and settle finished bouts."""
    from .predict import predict_card
    new_rows = []
    seen = set()
    for item in (previous, current):
        if item is None or item[0] is None:
            continue
        card, event, when = item
        card = card[card["odds_a"].notna() & card["odds_b"].notna()] if "odds_a" in card else card.iloc[0:0]
        key = (event, str(when), tuple(zip(card["fighter_a"], card["fighter_b"])))
        if card.empty or key in seen:
            continue
        seen.add(key)
        picks, _ = predict_card(card[["fighter_a", "fighter_b", "odds_a", "odds_b"]].reset_index(drop=True),
                                event, when)
        new_rows += tracking_mod.picks_to_rows(picks)
    before = {tracking_mod.row_key(r) for r in rows}
    merged = tracking_mod.merge(rows, new_rows)
    added = sum(1 for r in merged if tracking_mod.row_key(r) not in before)
    n = tracking_mod.settle(merged, fights)
    print(f"tracking: {added} new bouts logged, {n} results filled in")
    return merged


def _card_display_names(card: pd.DataFrame, fights, fighters, profiles) -> pd.DataFrame:
    """Swap each fighter's name for the exact spelling used in the Fighters sheet."""
    from .predict import resolve_card
    keep = [c for c in ("odds_source",) if c in card]
    extra = card[keep].reset_index(drop=True) if keep else None
    card = resolve_card(card.reset_index(drop=True), fights, fighters)
    if extra is not None:
        card[keep] = extra
    names = dict(zip(profiles["fighter_id"], profiles["name"]))
    for s in ("a", "b"):
        card[f"fighter_{s}"] = [names.get(i, n) for i, n in zip(card[f"id_{s}"], card[f"fighter_{s}"])]
    return card


# --------------------------------------------------------------------------- Fighters

def _fighter_table(fights, rounds, fighters, priors, as_of) -> pd.DataFrame:
    prof = features.current_profiles(fights, rounds, fighters, priors, as_of)
    prof = prof[prof["ufc_fights"] > 0]
    long = pd.concat([
        fights[["date", "id_a", "name_a", "division"]].set_axis(["date", "fighter_id", "name", "division"], axis=1),
        fights[["date", "id_b", "name_b", "division"]].set_axis(["date", "fighter_id", "name", "division"], axis=1),
    ]).sort_values("date")
    last = long.groupby("fighter_id").last()
    prof = prof.join(last[["name", "division"]], on="fighter_id")
    prof = prof.join(fighters.set_index("fighter_id")["dob"], on="fighter_id")
    # make display names unique (a few fighters share a name)
    dup = prof["name"].duplicated(keep=False)
    prof.loc[dup, "name"] = prof.loc[dup, "name"] + " (" + prof.loc[dup, "division"].str.title() + ")"
    dup = prof["name"].duplicated(keep=False)
    prof.loc[dup, "name"] = prof.loc[dup, "name"] + " #" + (prof[dup].groupby("name").cumcount() + 1).astype(str)
    return prof.sort_values("name", key=lambda s: s.str.lower()).reset_index(drop=True)


FIGHTER_COLS = [  # (header, profile column, number format)
    ("Fighter", "name", None), ("Last division", "division", None),
    ("UFC W", "ufc_wins", "0"), ("UFC L", "ufc_losses", "0"), ("UFC fights", "ufc_fights", "0"),
    ("UFC minutes", "ufc_minutes", "0.0"), ("SLpM", "slpm", "0.00"), ("Str Acc", "str_acc", "0.0%"),
    ("SApM", "sapm", "0.00"), ("Str Def", "str_def", "0.0%"), ("TD Avg", "td_avg", "0.00"),
    ("TD Acc", "td_acc", "0.0%"), ("TD Def", "td_def", "0.0%"), ("Sub Avg", "sub_avg", "0.00"),
    ("Ctrl min/15", "ctrl_avg", "0.00"), ("Elo", "elo", "0"), ("Date of birth", "dob", "yyyy-mm-dd"),
    ("Reach (in)", "reach", "0.0"), ("Last fight", "prev_date", "yyyy-mm-dd"),
]


def _fighters_sheet(ws, prof: pd.DataFrame, last_event) -> dict[str, str]:
    ws["A1"] = f"Every UFC fighter's stats going into their next fight (data up to {last_event.date()})"
    ws["A1"].font = _font(bold=True, size=12)
    ws["A2"] = ("Career rates are from UFC bouts only, pulled toward the league average when the sample is small. "
                "Regenerate this workbook (python3 -m fightiq workbook) after each event to refresh.")
    ws["A2"].font = _font(italic=True, color="595959")
    _header(ws, 4, [h for h, _, _ in FIGHTER_COLS])
    letters = {}
    for j, (h, col, fmt) in enumerate(FIGHTER_COLS, start=1):
        letters[col] = get_column_letter(j)
        for i, v in enumerate(prof[col].tolist(), start=5):
            if isinstance(v, float) and np.isnan(v):
                v = None
            elif isinstance(v, pd.Timestamp):
                v = None if pd.isna(v) else v.to_pydatetime()
            c = ws.cell(i, j, v)
            if fmt:
                c.number_format = fmt
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(FIGHTER_COLS))}{4 + len(prof)}"
    _widths(ws, {"A": 28, "B": 18})
    for j in range(3, len(FIGHTER_COLS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 11
    letters["_first"], letters["_last"] = 5, 4 + len(prof)
    return letters


# --------------------------------------------------------------------------- Model

def _model_sheet(ws, model: FightModel, priors: dict) -> dict[str, str]:
    """Writes every number the Predictor uses; returns their absolute cell addresses."""
    ws["A1"] = "FightIQ model parameters"
    ws["A1"].font = _font(bold=True, size=12)
    ws["A2"] = (f"Fitted on {model.meta['trained_fights']:,} UFC bouts "
                f"({model.meta['trained_from']} to {model.meta['trained_to']}) by python3 -m fightiq train. "
                "Values are written by the generator; change them only if you know why.")
    ws["A2"].font = _font(italic=True, color="595959")
    _header(ws, 4, ["Feature", "What it measures", "Coefficient"])
    cells = {}
    for i, (f, c) in enumerate(zip(model.features, model.coef), start=5):
        ws.cell(i, 1, f)
        ws.cell(i, 2, FEATURE_TEXT.get(f, f))
        k = ws.cell(i, 3, float(c))
        k.number_format = "0.0000"
        k.font = _font(color=BLUE)
        cells["coef_" + f] = f"Model!$C${i}"
    r = 5 + len(model.features) + 1
    ws.cell(r, 1, "P(A wins) = 1 / (1 + e^(−score)), score = Σ coefficient × feature. "
                  "No intercept: swapping the fighters flips the probability exactly.").font = _font(italic=True)

    bw = model.meta["blend_weights"]
    params = [
        ("blend_market", "Blend weight on logit(market probability)", bw["market_logit"], "0.000"),
        ("blend_model", "Blend weight on model score", bw["model_score"], "0.000"),
        ("threshold", "Bet when expected value is above", model.meta.get("bet_threshold", 0.05), "0.0%"),
        ("kelly", "Fraction of Kelly stake", 0.25, "0.00"),
        ("max_stake", "Maximum stake (% of bankroll)", 0.02, "0.0%"),
        ("elo_start", "Elo for a fighter with no UFC bouts", features.ELO_START, "0"),
        ("default_age", "Age used when date of birth is unknown", features.DEFAULT_AGE, "0.0"),
        ("default_layoff", "Days since last fight used for debutants", features.DEFAULT_LAYOFF_DAYS, "0"),
        ("prior_slpm", "League average SLpM (also SApM) - used for debutants", priors["slpm"], "0.00"),
        ("prior_acc", "League average strike accuracy - used for debutants", priors["str_acc"], "0.0%"),
        ("prior_td", "League average TD per 15 min - used for debutants", priors["td_avg"], "0.00"),
        ("prior_ctrl", "League average control min/15 - used for debutants", priors["ctrl_avg"], "0.00"),
    ]
    r += 2
    _header(ws, r, ["Setting", "Description", "Value"])
    for i, (key, desc, val, fmt) in enumerate(params, start=r + 1):
        ws.cell(i, 1, key)
        ws.cell(i, 2, desc)
        v = ws.cell(i, 3, float(val))
        v.number_format = fmt
        v.font = _font(color=BLUE)
        cells[key] = f"Model!$C${i}"
    r = r + len(params) + 2
    ws.cell(r, 1, "FightIQ P(A) with odds = 1 / (1 + e^(−z)), z = blend_market × ln(m / (1 − m)) + blend_model × score, "
                  "where m is the bookmaker probability with the margin removed: (1/oddsA) / (1/oddsA + 1/oddsB).").font = _font(italic=True)
    ws.cell(r + 1, 1, "Stake = min(kelly × (p × odds − 1) / (odds − 1), max_stake). "
                      "No bet on a bout with a UFC debutant.").font = _font(italic=True)
    _widths(ws, {"A": 16, "B": 58, "C": 12})
    return cells


# --------------------------------------------------------------------------- Predictor

# Per-side helper columns: (key, header, formula template). {m}=position in the Fighters
# list (0 when not found), {R:col}=that Fighters column's data range, {M:key}=Model cell.
SIDE_HELPERS = [
    ("slpm", "SLpM", "IF({m}>0,INDEX({R:slpm},{m}),{M:prior_slpm})"),
    ("sapm", "SApM", "IF({m}>0,INDEX({R:sapm},{m}),{M:prior_slpm})"),
    ("acc", "Str Acc", "IF({m}>0,INDEX({R:str_acc},{m}),{M:prior_acc})"),
    ("def", "Str Def", "IF({m}>0,INDEX({R:str_def},{m}),1-{M:prior_acc})"),
    ("td", "TD Avg", "IF({m}>0,INDEX({R:td_avg},{m}),{M:prior_td})"),
    ("ctrl", "Ctrl/15", "IF({m}>0,INDEX({R:ctrl_avg},{m}),{M:prior_ctrl})"),
    ("elo", "Elo", "IF({m}>0,INDEX({R:elo},{m}),{M:elo_start})"),
    ("fights", "UFC fights", "IF({m}>0,INDEX({R:ufc_fights},{m}),0)"),
    ("age", "Age", "IF({m}>0,IF(INDEX({R:dob},{m})=\"\",{M:default_age},"
                   "($B$3-INDEX({R:dob},{m}))/365.25),{M:default_age})"),
    ("reach", "Reach", "IF({m}>0,INDEX({R:reach},{m}),\"\")"),
    ("layoff", "Days off", "IF({m}>0,MAX($B$3-INDEX({R:prev_date},{m}),0),{M:default_layoff})"),
]

FEATURE_FORMULAS = {
    "slpm": "{A:slpm}-{B:slpm}",
    "exp_str_acc": "SQRT({A:acc}*(1-{B:def}))-SQRT({B:acc}*(1-{A:def}))",
    "sapm": "{A:sapm}-{B:sapm}",
    "td_avg": "{A:td}-{B:td}",
    "reach": "IF(OR({A:reach}=\"\",{B:reach}=\"\"),0,{A:reach}-{B:reach})",
    "age": "{A:age}-{B:age}",
    "elo": "({A:elo}-{B:elo})/100",
    "ctrl_avg": "{A:ctrl}-{B:ctrl}",
    "experience": "LN(1+{A:fights})-LN(1+{B:fights})",
    "layoff": "LN(1+{A:layoff})-LN(1+{B:layoff})",
    "age_old": "MAX({A:age}-33,0)-MAX({B:age}-33,0)",
    "debut": "({A:fights}=0)-({B:fights}=0)",
}

OUTPUT_HEADERS = ["Fighter A", "Fighter B", "Odds A (decimal)", "Odds B (decimal)", "Model P(A)", "Market P(A)",
                  "FightIQ P(A)", "Pick", "Pick prob.", "Bet", "Bet odds", "EV", "Stake (% bankroll)", "Notes"]


def _predictor_sheet(ws, model, cells, fcols, n_fighters, profiles, card=None, event_name="", event_date=None) -> None:
    first, last = fcols["_first"], fcols["_last"]
    ws["A1"] = f"FightIQ Predictor - {event_name}" if event_name else "FightIQ Predictor"
    ws["A1"].font = _font(bold=True, size=14)
    ws["A2"] = ("Type the fighters and decimal odds in the yellow cells (names must match the Fighters sheet; "
                "use the dropdown). Everything else calculates. Leave a row blank to skip it.")
    ws["A2"].font = _font(italic=True, color="595959")
    ws["A3"] = "Event date:"
    ws["A3"].font = _font(bold=True)
    ws["B3"] = pd.Timestamp(event_date).to_pydatetime() if event_date is not None else "=TODAY()"
    ws["B3"].number_format = "yyyy-mm-dd"
    ws["B3"].fill = INPUT_FILL
    ws["B3"].font = _font(color=BLUE, bold=True)
    ws["B3"].comment = Comment("Fight night. Used for age and days since last fight. Overwrite with the event date.", "FightIQ")

    hdr = 5
    _header(ws, hdr, OUTPUT_HEADERS)
    # helper block to the right
    col = len(OUTPUT_HEADERS) + 2
    helper_start = col
    ref = {}
    ws.cell(hdr - 1, col, "Calculation columns (don't edit)").font = _font(bold=True, color="595959")
    ref["mA"] = get_column_letter(col); col += 1
    ref["mB"] = get_column_letter(col); col += 1
    side_cols = {}
    for side in ("A", "B"):
        for key, _, _ in SIDE_HELPERS:
            side_cols[(side, key)] = get_column_letter(col)
            col += 1
    feat_cols = {}
    for f in model.features:
        feat_cols[f] = get_column_letter(col)
        col += 1
    ref["score"] = get_column_letter(col); col += 1
    ref["evA"] = get_column_letter(col); col += 1
    ref["evB"] = get_column_letter(col); col += 1
    helper_end = col - 1

    labels = ["Row in Fighters (A)", "Row in Fighters (B)"]
    labels += [f"{s}: {h}" for s in ("A", "B") for _, h, _ in SIDE_HELPERS]
    labels += [f"feature: {f}" for f in model.features] + ["score", "EV backing A", "EV backing B"]
    _header(ws, hdr, labels, start_col=helper_start, fill=SUB_FILL, color="000000")
    # coefficient row above the feature columns, linked to the Model sheet
    ws.cell(hdr - 1, helper_start + 2, "coefficients →").font = _font(italic=True, color="595959")
    for f in model.features:
        c = ws[f"{feat_cols[f]}{hdr - 1}"]
        c.value = f"={cells['coef_' + f]}"
        c.font = _font(color=GREEN)
        c.number_format = "0.0000"
    coef_range = f"${feat_cols[model.features[0]]}${hdr - 1}:${feat_cols[model.features[-1]]}${hdr - 1}"

    def fill(template: str, r: int, side: str | None = None) -> str:
        out = template
        if side:
            out = out.replace("{m}", f"${ref['m' + side]}{r}")
        for key, letter in fcols.items():
            if not key.startswith("_"):
                out = out.replace("{R:" + key + "}", f"Fighters!${letter}${first}:${letter}${last}")
        for key, addr in cells.items():
            out = out.replace("{M:" + key + "}", addr)
        for (s, key), letter in side_cols.items():
            out = out.replace("{" + s + ":" + key + "}", f"{letter}{r}")
        return out

    th, mk = cells["threshold"], None
    names_rng = f"Fighters!$A${first}:$A${last}"
    for r in range(hdr + 1, hdr + 1 + PREDICTOR_ROWS):
        for j in range(1, 5):
            c = ws.cell(r, j)
            c.fill = INPUT_FILL
            c.font = _font(color=BLUE)
            if j >= 3:
                c.number_format = "0.00"
        ws[f"{ref['mA']}{r}"] = f'=IFERROR(MATCH(TRIM(A{r}),{names_rng},0),0)'
        ws[f"{ref['mB']}{r}"] = f'=IFERROR(MATCH(TRIM(B{r}),{names_rng},0),0)'
        for side in ("A", "B"):
            for key, _, tmpl in SIDE_HELPERS:
                ws[f"{side_cols[(side, key)]}{r}"] = "=" + fill(tmpl, r, side)
        for f in model.features:
            ws[f"{feat_cols[f]}{r}"] = "=" + fill(FEATURE_FORMULAS[f], r)
        first_f, last_f = feat_cols[model.features[0]], feat_cols[model.features[-1]]
        ws[f"{ref['score']}{r}"] = f"=SUMPRODUCT({first_f}{r}:{last_f}{r},{coef_range})"
        blank = f'OR(A{r}="",B{r}="")'
        has_odds = f"AND(ISNUMBER(C{r}),ISNUMBER(D{r}))"
        ws[f"E{r}"] = f'=IF({blank},"",1/(1+EXP(-{ref["score"]}{r})))'
        ws[f"F{r}"] = f'=IF(OR({blank},NOT({has_odds})),"",(1/C{r})/(1/C{r}+1/D{r}))'
        ws[f"G{r}"] = (f'=IF({blank},"",IF(F{r}="",E{r},1/(1+EXP(-({cells["blend_market"]}*LN(F{r}/(1-F{r}))'
                       f'+{cells["blend_model"]}*{ref["score"]}{r})))))')
        ws[f"H{r}"] = f'=IF({blank},"",IF(G{r}>=0.5,A{r},B{r}))'
        ws[f"I{r}"] = f'=IF({blank},"",MAX(G{r},1-G{r}))'
        ws[f"{ref['evA']}{r}"] = f'=IF(F{r}="","",G{r}*C{r}-1)'
        ws[f"{ref['evB']}{r}"] = f'=IF(F{r}="","",(1-G{r})*D{r}-1)'
        evA, evB = f"{ref['evA']}{r}", f"{ref['evB']}{r}"
        fa, fb = f"{side_cols[('A', 'fights')]}{r}", f"{side_cols[('B', 'fights')]}{r}"
        ws[f"L{r}"] = f'=IF(F{r}="","",MAX({evA},{evB}))'
        ws[f"J{r}"] = f'=IF(L{r}="","",IF(AND(L{r}>{th},{fa}>0,{fb}>0),IF({evA}>={evB},A{r},B{r}),""))'
        ws[f"K{r}"] = f'=IF(J{r}="","",IF({evA}>={evB},C{r},D{r}))'
        p_side = f"IF({evA}>={evB},G{r},1-G{r})"
        ws[f"M{r}"] = (f'=IF(J{r}="","",MIN({cells["kelly"]}*({p_side}*K{r}-1)/(K{r}-1),{cells["max_stake"]}))')
        ws[f"N{r}"] = (f'=IF({blank},"",TRIM(IF(${ref["mA"]}{r}=0,A{r}&" not in Fighters (debut?). ","")'
                       f'&IF(${ref["mB"]}{r}=0,B{r}&" not in Fighters (debut?). ","")'
                       f'&IF(AND(L{r}<>"",L{r}>{th},OR({fa}=0,{fb}=0)),"No bet on debut bouts.","")))')
        for col_letter, fmt in (("E", "0.0%"), ("F", "0.0%"), ("G", "0.0%"), ("I", "0.0%"), ("K", "0.00"),
                                ("L", "0.0%"), ("M", "0.0%")):
            ws[f"{col_letter}{r}"].number_format = fmt
        for j in range(helper_start, helper_end + 1):
            ws.cell(r, j).fill = HELPER_FILL
            ws.cell(r, j).number_format = "0.000"
        for c in (ref["mA"], ref["mB"]):
            ws[f"{c}{r}"].number_format = "0"

    if card is not None:
        for i, row in card.head(PREDICTOR_ROWS).iterrows():
            r = hdr + 1 + i
            ws[f"A{r}"], ws[f"B{r}"] = row["fighter_a"], row["fighter_b"]
            for col, key in (("C", "odds_a"), ("D", "odds_b")):
                v = row.get(key)
                if v is not None and v == v:
                    ws[f"{col}{r}"] = float(v)
            source = row.get("odds_source")
            if isinstance(source, str) and source:
                ws[f"C{r}"].comment = Comment(f"Best price found: {source}", "FightIQ")
        pair = []
    else:
        # example bout: two experienced fighters from the data
        examples = profiles.sort_values("ufc_fights", ascending=False)
        recent = examples[examples["prev_date"] >= examples["prev_date"].max() - pd.Timedelta(days=200)]
        pair = recent.head(2)["name"].tolist()
    if len(pair) == 2:
        ws[f"A{hdr + 1}"], ws[f"B{hdr + 1}"] = pair
        ws[f"A{hdr + 1}"].comment = Comment("Example row (not a scheduled bout) - overwrite it with a real one and "
                                            "add decimal odds in columns C and D, e.g. 1.80 and 2.05.", "FightIQ")

    dv = DataValidation(type="list", formula1=f"={names_rng}", allow_blank=True, showErrorMessage=True,
                        errorStyle="warning", errorTitle="Fighter not found",
                        error="Not in the Fighters sheet. A UFC debutant gets league-average stats and no bet.")
    ws.add_data_validation(dv)
    dv.add(f"A{hdr + 1}:B{hdr + PREDICTOR_ROWS}")

    r = hdr + PREDICTOR_ROWS + 2
    notes = [
        "Legend: yellow cells with blue text are inputs. Odds are decimal (e.g. 1.80). For American odds: "
        "+150 → 2.50, −200 → 1.50.",
        "Model P(A): stats-only model. Market P(A): bookmaker odds with the margin removed. "
        "FightIQ P(A): the blend of both (use this one when you have odds).",
        "Bet: suggested when EV is above the Model sheet threshold and neither fighter is debuting. "
        "Stake is a quarter-Kelly % of bankroll, capped (see Model sheet).",
        "In the backtest the stats model alone did not beat the market; the blend did. Track every bet in the Tracking sheet.",
    ]
    for i, t in enumerate(notes):
        ws.cell(r + i, 1, t).font = _font(italic=True, color="595959")
    _widths(ws, {"A": 26, "B": 26, "C": 10, "D": 10, "E": 10, "F": 10, "G": 11, "H": 24, "I": 9, "J": 24,
                 "K": 9, "L": 8, "M": 10, "N": 46})
    for j in range(helper_start, helper_end + 1):
        ws.column_dimensions[get_column_letter(j)].width = 11
    ws.column_dimensions.group(get_column_letter(helper_start), get_column_letter(helper_end), hidden=False)
    ws.row_dimensions[hdr].height = 30
    ws.freeze_panes = f"C{hdr + 1}"


# --------------------------------------------------------------------------- Tracking

def _tracking_sheet(ws, existing_rows: list[list] | None = None) -> None:
    ws["A1"] = "FightIQ live tracking"
    ws["A1"].font = _font(bold=True, size=14)
    ws["A2"] = ("Filled in automatically: each weekly rebuild logs every Predictor bout that has odds, and fills in "
                "W/L once the result is in the data. You can also type rows yourself (W / L, D draw, NC no contest). "
                "Stake units: 1 unit = 1% of bankroll.")
    ws["A2"].font = _font(italic=True, color="595959")
    n_rows, top = 500, TRACKING_TOP
    existing_rows = existing_rows or []
    n_rows = max(n_rows, len(existing_rows) + 200)
    rng = lambda col: f"{col}{top}:{col}{top + n_rows - 1}"
    summary = [
        ("Picks (W-L)", f'=COUNTIF({rng("F")},"W")&"-"&COUNTIF({rng("F")},"L")', None),
        ("Pick hit rate", f'=IFERROR(COUNTIF({rng("F")},"W")/(COUNTIF({rng("F")},"W")+COUNTIF({rng("F")},"L")),"")', "0.0%"),
        ("Bets settled", f'=COUNT({rng("L")})', "0"),
        ("Units staked", f'=SUMPRODUCT(--ISNUMBER({rng("L")}),{rng("I")})', "0.00"),
        ("Profit (units)", f'=SUM({rng("L")})', "+0.00;-0.00;0.00"),
        ("ROI", f'=IFERROR(B7/B6,"")', "0.0%"),
    ]
    for i, (label, formula, fmt) in enumerate(summary, start=3):
        ws.cell(i, 1, label).font = _font(bold=True)
        c = ws.cell(i, 2, formula)
        if fmt:
            c.number_format = fmt
    # rows 3..8 hold the summary: B6 = units staked, B7 = profit
    headers = ["Date", "Event", "Fight", "Pick", "FightIQ prob.", "Pick result (W/L)", "Bet on", "Bet odds",
               "Stake (units)", "Bet result (W/L)", "Pick correct", "Bet profit (units)"]
    _header(ws, top - 1, headers)
    for r in range(top, top + n_rows):
        for col in "ABCDEFGHIJ":
            ws[f"{col}{r}"].fill = INPUT_FILL
            ws[f"{col}{r}"].font = _font(color=BLUE)
        ws[f"A{r}"].number_format = "yyyy-mm-dd"
        ws[f"E{r}"].number_format = "0.0%"
        ws[f"H{r}"].number_format = "0.00"
        ws[f"K{r}"] = f'=IF(F{r}="W",1,IF(F{r}="L",0,""))'
        ws[f"L{r}"] = (f'=IF(OR(G{r}="",J{r}=""),"",IF(J{r}="W",I{r}*(H{r}-1),IF(J{r}="L",-I{r},'
                       f'IF(OR(J{r}="D",J{r}="NC"),0,""))))')
        ws[f"L{r}"].number_format = "+0.00;-0.00;0.00"
    if existing_rows:
        for i, row in enumerate(existing_rows):
            for j, v in enumerate(row, start=1):
                ws.cell(top + i, j, v)
    else:
        example = [date(2026, 10, 3), "UFC 332 (example row)", "Figueiredo vs Talbott", "Payton Talbott", 0.66, "W",
                   "Payton Talbott", 1.60, 1, "W"]
        for j, v in enumerate(example, start=1):
            ws.cell(top, j, v)
        ws.cell(top, 2).comment = Comment("Example row showing the format - delete or overwrite it.", "FightIQ")
    _widths(ws, {"A": 12, "B": 26, "C": 32, "D": 22, "E": 11, "F": 11, "G": 22, "H": 9, "I": 9, "J": 11, "K": 9, "L": 12})
    ws.row_dimensions[top - 1].height = 30
    ws.freeze_panes = f"A{top}"


# --------------------------------------------------------------------------- Backtest

def _backtest_sheet(ws, P: pd.DataFrame) -> dict:
    mk = P[P["p_market"].notna() & P["p_blend"].notna()]
    nd = mk[(mk["ufc_fights_a"] > 0) & (mk["ufc_fights_b"] > 0)]
    ws["A1"] = "Walk-forward backtest"
    ws["A1"].font = _font(bold=True, size=14)
    ws["A2"] = ("Each year is predicted by a model trained only on earlier years, with every fighter stat computed "
                "from bouts before the fight. Odds: public dataset built from bestfightodds.com "
                "(github.com/shortlikeafox/ultimate_ufc_dataset); whether they are opening or closing lines is not documented.")
    ws["A2"].font = _font(italic=True, color="595959")

    r = 4
    ws.cell(r, 1, f"Prediction quality on the {len(mk):,} bouts with odds").font = _font(bold=True)
    _header(ws, r + 1, ["Model", "Bouts", "Accuracy", "Log loss", "Brier"])
    rows = [("Stats model only", "p_model"), ("Bookmaker odds (margin removed)", "p_market"), ("Blend (FightIQ)", "p_blend")]
    for i, (label, col) in enumerate(rows, start=r + 2):
        s = ev.scores(mk["result_a"], mk[col])
        for j, (v, fmt) in enumerate([(label, None), (s["n"], "#,##0"), (s["accuracy"], "0.0%"),
                                      (s["log_loss"], "0.0000"), (s["brier"], "0.0000")], start=1):
            c = ws.cell(i, j, v)
            if fmt:
                c.number_format = fmt
    ws.cell(r + 5, 1, "Lower log loss / Brier is better.").font = _font(italic=True, color="595959")

    r = 11
    ws.cell(r, 1, "Betting backtest - flat 1 unit, bouts without a UFC debutant").font = _font(bold=True)
    _header(ws, r + 1, ["Strategy", "Bets", "Profit (units)", "ROI"])
    i = r + 2
    main = None
    for col, label in (("p_model", "Stats model alone"), ("p_blend", "Blend")):
        for t in ev.BET_THRESHOLDS:
            sim = ev.betting_sim(nd, col, t)
            b = sim[sim["bet"]]
            for j, (v, fmt) in enumerate([(f"{label}, EV > {t:.0%}", None), (len(b), "#,##0"),
                                          (b["pnl"].sum(), "+0.0;-0.0"), (b["pnl"].mean(), "0.0%")], start=1):
                c = ws.cell(i, j, float(v) if isinstance(v, (np.floating, float)) else v)
                if fmt:
                    c.number_format = fmt
            if col == "p_blend" and t == 0.05:
                main = (sim, b)
                for j in range(1, 5):
                    ws.cell(i, j).font = _font(bold=True)
            i += 1

    sim, b = main
    lo, hi = ev.event_bootstrap_roi(sim)
    ws.cell(i + 1, 1, f"Rule used by the Predictor (blend, EV > 5%): 95% interval for ROI {lo:+.1%} to {hi:+.1%}").font = _font(italic=True)

    r = i + 3
    ws.cell(r, 1, "Blend, EV > 5%: by year").font = _font(bold=True)
    _header(ws, r + 1, ["Year", "Bets", "Profit (units)", "ROI", "Cumulative units"])
    per = b.groupby("year")["pnl"].agg(["size", "sum", "mean"])
    cum = 0.0
    first_row = r + 2
    for k, (y, row) in enumerate(per.iterrows()):
        cum += row["sum"]
        vals = [(str(int(y)), None), (int(row["size"]), "0"), (float(row["sum"]), "+0.0;-0.0"),
                (float(row["mean"]), "0.0%"), (cum, "+0.0;-0.0")]
        for j, (v, fmt) in enumerate(vals, start=1):
            c = ws.cell(first_row + k, j, v)
            if fmt:
                c.number_format = fmt
    last_row = first_row + len(per) - 1
    chart = BarChart()
    chart.title = "Blend bets: profit per year (units)"
    chart.y_axis.title = "Units"
    chart.add_data(Reference(ws, min_col=3, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first_row, max_row=last_row))
    chart.legend = None
    chart.height, chart.width = 7.5, 16
    ws.add_chart(chart, "G4")

    r = last_row + 3
    ws.cell(r, 1, "Calibration of the blend (favoured side)").font = _font(bold=True)
    _header(ws, r + 1, ["Predicted", "Bouts", "Avg predicted", "Actually won"])
    cal = ev.calibration(mk["result_a"], mk["p_blend"])
    for k, row in enumerate(cal.itertuples(), start=r + 2):
        for j, (v, fmt) in enumerate([(row.predicted, None), (row.fights, "#,##0"), (row.avg_predicted, "0.0%"),
                                      (row.actual_win_rate, "0.0%")], start=1):
            c = ws.cell(k, j, v)
            if fmt:
                c.number_format = fmt
    _widths(ws, {"A": 40, "B": 10, "C": 14, "D": 12, "E": 16})
    return {"bets": len(b), "roi": float(b["pnl"].mean()), "profit": float(b["pnl"].sum()), "ci": (lo, hi),
            "acc_model": ev.scores(mk["result_a"], mk["p_model"])["accuracy"],
            "acc_market": ev.scores(mk["result_a"], mk["p_market"])["accuracy"],
            "acc_blend": ev.scores(mk["result_a"], mk["p_blend"])["accuracy"],
            "years": f"{int(P['year'].min())}–{int(P['year'].max())}"}


# --------------------------------------------------------------------------- Fight Database

def _database_sheet(ws, X: pd.DataFrame, P: pd.DataFrame, model: FightModel) -> None:
    d = X[X["date"] >= "2001-01-01"].copy()
    oos = P.set_index("fight_id")[["p_model", "p_blend"]]
    d = d.join(oos, on="fight_id")
    in_sample = model.predict_proba(d)
    d["p_out_of_sample"] = d["p_model"].notna()
    d["p_model"] = d["p_model"].fillna(pd.Series(in_sample, index=d.index))
    d["winner"] = np.where(d["result_a"] == 1, d["name_a"], np.where(d["result_a"] == 0, d["name_b"],
                           np.where(d["result_a"] == 0.5, "Draw", "No contest")))
    stat_cols = [("ufc_wins", "UFC W", "0"), ("ufc_losses", "UFC L", "0"), ("slpm", "SLpM", "0.00"),
                 ("str_acc", "Str Acc", "0.0%"), ("sapm", "SApM", "0.00"), ("str_def", "Str Def", "0.0%"),
                 ("td_avg", "TD Avg", "0.00"), ("td_acc", "TD Acc", "0.0%"), ("td_def", "TD Def", "0.0%"),
                 ("sub_avg", "Sub Avg", "0.00"), ("ctrl_avg", "Ctrl/15", "0.00"), ("elo", "Elo", "0"),
                 ("age", "Age", "0.0"), ("reach", "Reach", "0")]
    columns = [("date", "Date", "yyyy-mm-dd"), ("event", "Event", None), ("division", "Division", None),
               ("name_a", "Fighter A", None), ("name_b", "Fighter B", None), ("winner", "Winner", None),
               ("method", "Method", None)]
    columns += [(f"{c}_a", f"A {h}", fmt) for c, h, fmt in stat_cols]
    columns += [(f"{c}_b", f"B {h}", fmt) for c, h, fmt in stat_cols]
    columns += [("f_" + f, f"feature: {f}", "0.000") for f in model.features]
    columns += [("p_model", "Model P(A)", "0.0%"), ("p_out_of_sample", "Out of sample?", None),
                ("odds_a", "Odds A", "0.00"), ("odds_b", "Odds B", "0.00"), ("p_market", "Market P(A)", "0.0%"),
                ("p_blend", "Blend P(A)", "0.0%")]
    ws["A1"] = ("Every UFC bout since 2001 with both fighters' stats as they stood BEFORE the bout (no look-ahead). "
                "Model P(A) is out of sample from 2015 (walk-forward); earlier rows use the final model.")
    ws["A1"].font = _font(bold=True)
    _header(ws, 2, [h for _, h, _ in columns])
    d = d.sort_values("date", ascending=False)
    for i, rec in enumerate(d[[c for c, _, _ in columns]].itertuples(index=False), start=3):
        for j, v in enumerate(rec, start=1):
            if isinstance(v, float) and np.isnan(v):
                continue
            if isinstance(v, (np.bool_, bool)):
                v = "Yes" if v else "No"
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif isinstance(v, np.generic):
                v = v.item()
            ws.cell(i, j, v)
    for j, (_, _, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = 11
        if fmt:
            for row in ws.iter_rows(min_row=3, max_row=2 + len(d), min_col=j, max_col=j):
                row[0].number_format = fmt
    _widths(ws, {"B": 30, "C": 16, "D": 22, "E": 22, "F": 22, "G": 18})
    ws.row_dimensions[2].height = 30
    ws.freeze_panes = "F3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{2 + len(d)}"


# --------------------------------------------------------------------------- Guide

def _guide_sheet(ws, model, s: dict, last_event, n_fights: int) -> None:
    ws["A1"] = "FightIQ - UFC Prediction Model"
    ws["A1"].font = _font(bold=True, size=16)
    ws["A2"] = f"Built {date.today().isoformat()}  ·  data up to {last_event.date()}  ·  {n_fights:,} UFC bouts since UFC 1"
    ws["A2"].font = _font(italic=True, color="595959")
    sections = [
        ("What it does", [
            "Predicts the win probability for each fighter in a UFC bout, then blends it with the betting odds to find "
            "bets with positive expected value.",
            "Every stat is what was known BEFORE each fight: career numbers are rebuilt from round-by-round ufcstats data "
            "using only earlier bouts. The old model used today's career stats for past fights, which leaked the result.",
        ]),
        ("How to use it each week", [
            "1. Close this file, then in Terminal run:  python3 -m fightiq workbook --upcoming  - it refreshes the stats, "
            "fills the Predictor with the next UFC card and the best current odds, logs the picks into Tracking, and "
            "fills in W/L for last week's bouts.",
            "2. Open the file. Type odds for any bout left blank in columns C and D (hover an odds cell to see the bookmaker).",
            "3. Read FightIQ P(A) (chance Fighter A wins), the Pick, and the Bet column with its suggested stake.",
            "4. If you typed or changed odds, save, close, and run  python3 -m fightiq workbook  to log them into Tracking. "
            "To update odds nearer fight time:  python3 -m fightiq workbook --refresh-odds",
            "5. After the event, the next weekly run fills in the results. Tracking rows are kept on every rebuild "
            "(a backup of the old file goes to output/backup).",
            "Typing fighters by hand: capitals don't matter (jim miller = Jim Miller), but the spelling must match the "
            "Fighters sheet - plain letters, no accents. The Notes column flags names it can't find.",
        ]),
        ("How it works", [
            "Features (fighter A minus fighter B): " + ", ".join(model.features) + ". See the Model sheet.",
            "Logistic regression trained on every bout from both corners with no intercept, so fighter order can't change "
            "the result (no red-corner bias).",
            "With odds: z = a × logit(bookmaker probability) + b × model score. The bookmaker margin is removed first.",
            "Bet rule: EV above the threshold, never on a UFC debut bout, quarter-Kelly stake capped at 2% of bankroll.",
        ]),
        (f"Backtest ({s['years']}, trained only on earlier years)", [
            f"Stats model alone picks {s['acc_model']:.1%} of winners; bookmaker favourites win {s['acc_market']:.1%}; "
            f"the blend picks {s['acc_blend']:.1%}.",
            f"Bet rule: {s['bets']} bets, {s['profit']:+.1f} units, ROI {s['roi']:+.1%} "
            f"(95% interval {s['ci'][0]:+.1%} to {s['ci'][1]:+.1%}).",
            "Caveat: it is not documented whether the historical odds are opening or closing prices. Opening lines would "
            "overstate the ROI. Treat it as promising, track live, and size small.",
        ]),
        ("Limitations", [
            "UFC debutants have no UFC stats: they get league-average numbers and are never bet.",
            "No injuries, weight cuts, short-notice changes or camp news. Late-notice fights deserve extra caution.",
            "The stats model alone is weaker than the betting market. The value comes from combining the two.",
        ]),
        ("Old model (2025 spreadsheet)", [
            "Live record 130-79 (62.2%), −15.7 units at 1 unit flat (−7.5% ROI). Its 74.5% backtest came from look-ahead "
            "stats and a test set where fighter 1 won ~90% of bouts. Details: reports/original_model_review.md.",
        ]),
        ("Colour legend", [
            "Yellow cell, blue text = input you fill in.  Black text = formula.  Green text = linked from another sheet.",
        ]),
    ]
    r = 4
    for title, lines in sections:
        c = ws.cell(r, 1, title)
        c.font = _font(bold=True, size=11, color="FFFFFF")
        c.fill = HEADER_FILL
        r += 1
        for line in lines:
            c = ws.cell(r, 1, line)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * max(1, int(len(line) / 120) + 1)
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 130
